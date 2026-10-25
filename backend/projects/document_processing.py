import csv
import io
import re
from pathlib import Path


MAX_PDF_PAGES = 75
MAX_WORKSHEETS = 30
MAX_SPREADSHEET_ROWS = 10000
MAX_EXTRACTED_CHARACTERS = 120000


class DocumentProcessingError(ValueError):
    pass


def _bounded(text):
    normalized = re.sub(r"\x00", "", str(text or "")).strip()
    return normalized[:MAX_EXTRACTED_CHARACTERS]


def _append_bounded(lines, value, used):
    if used >= MAX_EXTRACTED_CHARACTERS:
        return used, True
    text = re.sub(r"\x00", "", str(value or "")).strip()
    if not text:
        return used, False
    separator = 1 if lines else 0
    remaining = MAX_EXTRACTED_CHARACTERS - used - separator
    if remaining <= 0:
        return used, True
    lines.append(text[:remaining])
    used += separator + min(len(text), remaining)
    return used, used >= MAX_EXTRACTED_CHARACTERS


def _extract_pdf(data):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentProcessingError("PDF processing support is not installed.") from exc
    try:
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentProcessingError("Encrypted PDF files are not supported.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentProcessingError(f"PDF files are limited to {MAX_PDF_PAGES} pages.")
        native_lines = []
        used = 0
        for page in reader.pages:
            used, full = _append_bounded(native_lines, page.extract_text() or "", used)
            if full:
                break
        native = "\n".join(native_lines)
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("The PDF is corrupt or could not be read.") from exc
    if native.strip():
        return _bounded(native), {"method": "native_pdf", "pages": len(reader.pages)}
    try:
        import pytesseract
        from pdf2image import convert_from_bytes

        ocr_lines = []
        used = 0
        processed_pages = 0
        for page_number in range(1, len(reader.pages) + 1):
            images = convert_from_bytes(data, dpi=160, first_page=page_number, last_page=page_number)
            if not images:
                continue
            processed_pages += 1
            used, full = _append_bounded(ocr_lines, pytesseract.image_to_string(images[0]), used)
            if full:
                break
        text = "\n".join(ocr_lines)
    except Exception as exc:
        raise DocumentProcessingError("The scanned PDF could not be processed by OCR.") from exc
    if not text.strip():
        raise DocumentProcessingError("No readable text was found in the PDF.")
    return _bounded(text), {"method": "pdf_ocr", "pages": processed_pages}


def _extract_docx(data):
    try:
        from docx import Document

        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentProcessingError("The DOCX file is corrupt or could not be read.") from exc
    lines = []
    used = 0
    for paragraph in document.paragraphs:
        used, full = _append_bounded(lines, paragraph.text, used)
        if full:
            break
    for table in document.tables:
        if used >= MAX_EXTRACTED_CHARACTERS:
            break
        for row in table.rows:
            used, full = _append_bounded(lines, " | ".join(cell.text.strip() for cell in row.cells), used)
            if full:
                break
    text = _bounded("\n".join(lines))
    if not text:
        raise DocumentProcessingError("No readable text was found in the DOCX file.")
    return text, {"method": "docx", "tables": len(document.tables)}


def _extract_xlsx(data):
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentProcessingError("The XLSX file is corrupt or could not be read.") from exc
    if len(workbook.sheetnames) > MAX_WORKSHEETS:
        raise DocumentProcessingError(f"Spreadsheets are limited to {MAX_WORKSHEETS} worksheets.")
    lines = []
    row_count = 0
    used = 0
    for worksheet in workbook.worksheets:
        used, full = _append_bounded(lines, f"Worksheet: {worksheet.title}", used)
        if full:
            break
        for row in worksheet.iter_rows(values_only=True):
            row_count += 1
            if row_count > MAX_SPREADSHEET_ROWS:
                raise DocumentProcessingError(f"Spreadsheets are limited to {MAX_SPREADSHEET_ROWS} rows.")
            values = [str(value) for value in row if value not in (None, "")]
            if values:
                used, full = _append_bounded(lines, " | ".join(values), used)
            if full:
                break
    text = _bounded("\n".join(lines))
    if not text:
        raise DocumentProcessingError("No readable values were found in the XLSX file.")
    return text, {"method": "xlsx", "worksheets": len(workbook.sheetnames), "rows": row_count}


def _extract_csv(data):
    decoded = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise DocumentProcessingError("The CSV encoding is not supported.")
    try:
        rows = []
        used = 0
        for index, row in enumerate(csv.reader(io.StringIO(decoded)), start=1):
            if index > MAX_SPREADSHEET_ROWS:
                raise DocumentProcessingError(f"CSV files are limited to {MAX_SPREADSHEET_ROWS} rows.")
            used, full = _append_bounded(rows, " | ".join(row), used)
            if full:
                break
    except csv.Error as exc:
        raise DocumentProcessingError("The CSV file is corrupt or could not be read.") from exc
    text = _bounded("\n".join(rows))
    if not text:
        raise DocumentProcessingError("No readable values were found in the CSV file.")
    return text, {"method": "csv", "rows": len(rows)}


def _extract_image(data):
    try:
        import pytesseract
        from PIL import Image

        image = Image.open(io.BytesIO(data))
        image.verify()
        image = Image.open(io.BytesIO(data))
        image.load()
        text = pytesseract.image_to_string(image)
    except Exception as exc:
        raise DocumentProcessingError("The image is corrupt or could not be processed by OCR.") from exc
    if not text.strip():
        raise DocumentProcessingError("No readable text was found in the image.")
    return _bounded(text), {"method": "image_ocr", "width": image.width, "height": image.height}


def extract_document_text(document):
    data = bytes(document.content)
    extension = str(document.extension or Path(document.original_filename).suffix).lower().lstrip(".")
    if extension == "pdf":
        return _extract_pdf(data)
    if extension == "docx":
        return _extract_docx(data)
    if extension == "xlsx":
        return _extract_xlsx(data)
    if extension == "csv":
        return _extract_csv(data)
    if extension in {"png", "jpg", "jpeg"}:
        return _extract_image(data)
    raise DocumentProcessingError("This document format is not supported.")
